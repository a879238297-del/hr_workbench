"""
HR 智能工作台 · Flask 后端
启动: python app.py
访问: http://localhost:8080
"""
import io
import json
import os
import shutil
import sqlite3
import uuid
from datetime import datetime, timedelta, timezone
from functools import wraps
from pathlib import Path

import jwt
import openpyxl
from flask import Flask, g, jsonify, request, send_file, send_from_directory

from reconcile import reconcile
from contract_alert import analyze as contract_analyze
from roster_check import validate as roster_validate
from attendance import analyze as attendance_analyze
from social_ins import reconcile as social_reconcile
from cost_analysis import analyze as cost_analyze

# ── 配置 ────────────────────────────────────────────────────────────────────
BASE_DIR = Path(__file__).parent
DATA_DIR = BASE_DIR / "data"
UPLOAD_DIR = BASE_DIR / "uploads"
WEB_DIR = BASE_DIR / "web"

DATA_DIR.mkdir(exist_ok=True)
UPLOAD_DIR.mkdir(exist_ok=True)

DB_PATH = DATA_DIR / "hr_workbench.db"
SECRET_KEY = os.getenv("SECRET_KEY", "dev-secret-change-in-production")
JWT_EXPIRE_HOURS = 8
MAX_UPLOAD_BYTES = 50 * 1024 * 1024  # 50 MB

app = Flask(__name__, static_folder=None)
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_BYTES

# ── 数据库 ───────────────────────────────────────────────────────────────────
SCHEMA = """
CREATE TABLE IF NOT EXISTS app_user (
    id INTEGER PRIMARY KEY,
    username TEXT NOT NULL UNIQUE,
    password TEXT NOT NULL,
    display_name TEXT NOT NULL,
    is_active INTEGER NOT NULL DEFAULT 1
);

CREATE TABLE IF NOT EXISTS job (
    id TEXT PRIMARY KEY,
    module_key TEXT NOT NULL DEFAULT 'payroll_recon',
    period TEXT,
    status TEXT NOT NULL DEFAULT 'done',
    summary TEXT,
    created_by TEXT NOT NULL,
    created_at TEXT NOT NULL,
    finished_at TEXT
);

CREATE TABLE IF NOT EXISTS job_file (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    job_id TEXT NOT NULL,
    side TEXT NOT NULL,
    original_name TEXT NOT NULL,
    sha256_before TEXT NOT NULL,
    sha256_after TEXT NOT NULL,
    readonly_verified INTEGER NOT NULL
);

CREATE TABLE IF NOT EXISTS job_row (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    job_id TEXT NOT NULL,
    row_index INTEGER NOT NULL,
    match_key TEXT,
    category TEXT NOT NULL,
    payload TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_job_row ON job_row(job_id, category);

CREATE TABLE IF NOT EXISTS audit_log (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    username TEXT NOT NULL,
    action TEXT NOT NULL,
    target_id TEXT,
    created_at TEXT NOT NULL
);
"""

SEED_USERS = [
    ("hr01",    "hr123456",    "李静（HR）"),
    ("fin01",   "fin123456",   "王芳（财务）"),
    ("admin",   "admin888",    "陈工（管理员）"),
]


def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA journal_mode=WAL")
        g.db.execute("PRAGMA foreign_keys=ON")
    return g.db


@app.teardown_appcontext
def close_db(_):
    db = g.pop("db", None)
    if db:
        db.close()


def init_db():
    con = sqlite3.connect(DB_PATH)
    con.executescript(SCHEMA)
    for username, password, display_name in SEED_USERS:
        con.execute(
            "INSERT OR IGNORE INTO app_user(username, password, display_name) VALUES(?,?,?)",
            (username, password, display_name),
        )
    con.commit()
    con.close()


# ── JWT ─────────────────────────────────────────────────────────────────────
def make_token(username: str) -> str:
    payload = {
        "sub": username,
        "exp": datetime.now(timezone.utc) + timedelta(hours=JWT_EXPIRE_HOURS),
    }
    return jwt.encode(payload, SECRET_KEY, algorithm="HS256")


def verify_token(token: str) -> str | None:
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=["HS256"])
        return payload["sub"]
    except jwt.PyJWTError:
        return None


def require_auth(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        auth = request.headers.get("Authorization", "")
        if not auth.startswith("Bearer "):
            return jsonify({"code": "UNAUTHENTICATED", "message": "请先登录"}), 401
        username = verify_token(auth[7:])
        if not username:
            return jsonify({"code": "UNAUTHENTICATED", "message": "Token 已过期，请重新登录"}), 401
        g.current_user = username
        return f(*args, **kwargs)
    return wrapper


# ── 工具 ─────────────────────────────────────────────────────────────────────
def now_iso() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def err(code: str, message: str, status: int = 400, detail: dict | None = None):
    body = {"code": code, "message": message}
    if detail:
        body["detail"] = detail
    return jsonify(body), status


# ── 路由：静态文件 ───────────────────────────────────────────────────────────
@app.route("/", defaults={"path": ""})
@app.route("/<path:path>")
def serve_static(path):
    if path.startswith("api/"):
        return err("NOT_FOUND", "接口不存在", 404)
    target = WEB_DIR / (path or "index.html")
    if target.is_file():
        return send_from_directory(WEB_DIR, path or "index.html")
    return send_from_directory(WEB_DIR, "index.html")


# ── 路由：认证 ───────────────────────────────────────────────────────────────
@app.route("/api/auth/login", methods=["POST"])
def login():
    body = request.get_json(force=True, silent=True) or {}
    username = str(body.get("username", "")).strip()
    password = str(body.get("password", "")).strip()
    if not username or not password:
        return err("INVALID_PARAMS", "用户名和密码不能为空")

    db = get_db()
    row = db.execute(
        "SELECT * FROM app_user WHERE username=? AND is_active=1", (username,)
    ).fetchone()
    if not row or row["password"] != password:
        return err("AUTH_FAILED", "用户名或密码错误", 401)

    token = make_token(username)
    db.execute(
        "INSERT INTO audit_log(username, action, created_at) VALUES(?,?,?)",
        (username, "login", now_iso()),
    )
    db.commit()
    return jsonify({
        "access_token": token,
        "token_type": "bearer",
        "expires_in": JWT_EXPIRE_HOURS * 3600,
        "username": username,
        "display_name": row["display_name"],
    })


# ── 路由：模块列表（前端导航用） ────────────────────────────────────────────
MODULES = [
    {"key": "payroll_recon", "name": "薪资对账", "phase": 1, "enabled": True},
    {"key": "contract_alert", "name": "合同到期预警", "phase": 1, "enabled": True},
    {"key": "roster_check", "name": "花名册校验", "phase": 1, "enabled": True},
    {"key": "attendance", "name": "考勤统计", "phase": 1, "enabled": True},
    {"key": "social_ins", "name": "社保公积金核对", "phase": 1, "enabled": True},
    {"key": "tax_check", "name": "个税代扣核对", "phase": 2, "enabled": False},
    {"key": "recruit", "name": "招聘漏斗分析", "phase": 2, "enabled": False},
    {"key": "cost_analysis", "name": "人员成本分析", "phase": 1, "enabled": True},
    {"key": "performance", "name": "绩效考核汇总", "phase": 2, "enabled": False},
    {"key": "training", "name": "培训学时统计", "phase": 2, "enabled": False},
    {"key": "turnover", "name": "离职率分析", "phase": 2, "enabled": False},
    {"key": "overtime", "name": "加班费核算", "phase": 2, "enabled": False},
]


@app.route("/api/modules")
@require_auth
def list_modules():
    return jsonify(MODULES)


# ── 路由：作业 ───────────────────────────────────────────────────────────────
@app.route("/api/jobs", methods=["POST"])
@require_auth
def create_job():
    payroll_file = request.files.get("payroll_file")
    bank_file = request.files.get("bank_file")
    period = request.form.get("period", "").strip()

    if not payroll_file or not bank_file:
        return err("MISSING_FILE", "请同时上传工资表和银行代发明细")

    for f in [payroll_file, bank_file]:
        if not f.filename.lower().endswith((".xlsx", ".xls")):
            return err("FILE_TYPE_MISMATCH", f"仅支持 Excel 文件（.xlsx / .xls），收到：{f.filename}")

    job_id = str(uuid.uuid4())
    job_dir = UPLOAD_DIR / job_id
    job_dir.mkdir()

    payroll_path = job_dir / f"payroll_{payroll_file.filename}"
    bank_path = job_dir / f"bank_{bank_file.filename}"
    payroll_file.save(payroll_path)
    bank_file.save(bank_path)

    try:
        result = reconcile(payroll_path, bank_path)
    except ValueError as e:
        parts = str(e).split("|")
        code = parts[0] if len(parts) > 1 else "RECONCILE_ERROR"
        msg = parts[1] if len(parts) > 1 else str(e)
        detail = {}
        for p in parts[2:]:
            if "=" in p:
                k, v = p.split("=", 1)
                try:
                    detail[k] = json.loads(v)
                except Exception:
                    detail[k] = v
        shutil.rmtree(job_dir, ignore_errors=True)
        return err(code, msg, 422, detail or None)

    db = get_db()
    created_at = now_iso()
    db.execute(
        "INSERT INTO job(id, period, summary, created_by, created_at, finished_at)"
        " VALUES(?,?,?,?,?,?)",
        (job_id, period or None, json.dumps(result["summary"], ensure_ascii=False),
         g.current_user, created_at, created_at),
    )
    for f_info in result["files"]:
        f_meta = payroll_file if f_info["side"] == "payroll" else bank_file
        db.execute(
            "INSERT INTO job_file(job_id, side, original_name, sha256_before, sha256_after, readonly_verified)"
            " VALUES(?,?,?,?,?,?)",
            (job_id, f_info["side"], f_meta.filename,
             f_info["sha256_before"], f_info["sha256_after"],
             1 if f_info["readonly_verified"] else 0),
        )
    for i, row in enumerate(result["rows"]):
        db.execute(
            "INSERT INTO job_row(job_id, row_index, match_key, category, payload)"
            " VALUES(?,?,?,?,?)",
            (job_id, i, row.get("match_key"), row["category"],
             json.dumps(row, ensure_ascii=False)),
        )
    db.execute(
        "INSERT INTO audit_log(username, action, target_id, created_at) VALUES(?,?,?,?)",
        (g.current_user, "create_job", job_id, created_at),
    )
    db.commit()

    return jsonify(_job_response(job_id, period, result["summary"], created_at, result["files"],
                                 payroll_file.filename, bank_file.filename)), 201


@app.route("/api/contract-alert/jobs", methods=["POST"])
@require_auth
def create_contract_alert_job():
    contract_file = request.files.get("contract_file")
    if not contract_file:
        return err("MISSING_FILE", "请上传合同台账文件")
    if not contract_file.filename.lower().endswith((".xlsx", ".xls")):
        return err("FILE_TYPE_MISMATCH", f"仅支持 Excel 文件（.xlsx / .xls），收到：{contract_file.filename}")

    job_id = str(uuid.uuid4())
    job_dir = UPLOAD_DIR / job_id
    job_dir.mkdir()
    file_path = job_dir / f"contract_{contract_file.filename}"
    contract_file.save(file_path)

    try:
        result = contract_analyze(file_path)
    except ValueError as e:
        parts = str(e).split("|")
        code = parts[0] if len(parts) > 1 else "ANALYZE_ERROR"
        msg = parts[1] if len(parts) > 1 else str(e)
        detail = {}
        for p in parts[2:]:
            if "=" in p:
                k, v = p.split("=", 1)
                try:
                    detail[k] = json.loads(v)
                except Exception:
                    detail[k] = v
        shutil.rmtree(job_dir, ignore_errors=True)
        return err(code, msg, 422, detail or None)

    db = get_db()
    created_at = now_iso()
    db.execute(
        "INSERT INTO job(id, module_key, period, summary, created_by, created_at, finished_at)"
        " VALUES(?,?,?,?,?,?,?)",
        (job_id, "contract_alert", None,
         json.dumps(result["summary"], ensure_ascii=False),
         g.current_user, created_at, created_at),
    )
    for f_info in result["files"]:
        db.execute(
            "INSERT INTO job_file(job_id, side, original_name, sha256_before, sha256_after, readonly_verified)"
            " VALUES(?,?,?,?,?,?)",
            (job_id, f_info["side"], contract_file.filename,
             f_info["sha256_before"], f_info["sha256_after"],
             1 if f_info["readonly_verified"] else 0),
        )
    for i, row in enumerate(result["rows"]):
        db.execute(
            "INSERT INTO job_row(job_id, row_index, match_key, category, payload)"
            " VALUES(?,?,?,?,?)",
            (job_id, i, row.get("match_key"), row["category"],
             json.dumps(row, ensure_ascii=False)),
        )
    db.execute(
        "INSERT INTO audit_log(username, action, target_id, created_at) VALUES(?,?,?,?)",
        (g.current_user, "create_contract_alert_job", job_id, created_at),
    )
    db.commit()

    return jsonify({
        "id": job_id,
        "module_key": "contract_alert",
        "status": "done",
        "summary": result["summary"],
        "created_by": g.current_user,
        "created_at": created_at,
        "files": [
            {
                "side": f["side"],
                "original_name": contract_file.filename,
                "sha256_before": f["sha256_before"],
                "sha256_after": f["sha256_after"],
                "readonly_verified": f["readonly_verified"],
            }
            for f in result["files"]
        ],
    }), 201


# ── 花名册校验 ─────────────────────────────────────────────────────────────
@app.route("/api/roster-check/jobs", methods=["POST"])
@require_auth
def create_roster_check_job():
    roster_file = request.files.get("roster_file")
    if not roster_file:
        return err("MISSING_FILE", "请上传花名册文件")
    if not roster_file.filename.lower().endswith((".xlsx", ".xls")):
        return err("FILE_TYPE_MISMATCH", f"仅支持 Excel 文件（.xlsx / .xls），收到：{roster_file.filename}")

    job_id = str(uuid.uuid4())
    job_dir = UPLOAD_DIR / job_id
    job_dir.mkdir(parents=True, exist_ok=True)
    file_path = job_dir / roster_file.filename
    roster_file.save(str(file_path))

    try:
        result = roster_validate(file_path)
    except ValueError as exc:
        parts = str(exc).split("|")
        code = parts[0] if len(parts) >= 2 else "VALIDATE_ERROR"
        msg = parts[1] if len(parts) >= 2 else str(exc)
        return err(code, msg, 422)

    db = get_db()
    created_at = datetime.now(tz=timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
    db.execute(
        "INSERT INTO job(id, module_key, period, summary, created_by, created_at, finished_at)"
        " VALUES(?,?,?,?,?,?,?)",
        (job_id, "roster_check", None,
         json.dumps(result["summary"], ensure_ascii=False),
         g.current_user, created_at, created_at),
    )
    for f_info in result["files"]:
        db.execute(
            "INSERT INTO job_file(job_id, side, original_name, sha256_before, sha256_after, readonly_verified)"
            " VALUES(?,?,?,?,?,?)",
            (job_id, f_info["side"], roster_file.filename,
             f_info["sha256_before"], f_info["sha256_after"],
             1 if f_info["readonly_verified"] else 0),
        )
    for i, row in enumerate(result["rows"]):
        db.execute(
            "INSERT INTO job_row(job_id, row_index, match_key, category, payload)"
            " VALUES(?,?,?,?,?)",
            (job_id, i, row.get("match_key"), row["category"],
             json.dumps(row, ensure_ascii=False)),
        )
    db.execute(
        "INSERT INTO audit_log(username, action, target_id, created_at) VALUES(?,?,?,?)",
        (g.current_user, "create_roster_check_job", job_id, created_at),
    )
    db.commit()

    return jsonify({
        "id": job_id,
        "module_key": "roster_check",
        "status": "done",
        "summary": result["summary"],
        "created_by": g.current_user,
        "created_at": created_at,
        "files": [
            {
                "side": f["side"],
                "original_name": roster_file.filename,
                "sha256_before": f["sha256_before"],
                "sha256_after": f["sha256_after"],
                "readonly_verified": f["readonly_verified"],
            }
            for f in result["files"]
        ],
    }), 201


# ── 考勤统计 ─────────────────────────────────────────────────────────────
@app.route("/api/attendance/jobs", methods=["POST"])
@require_auth
def create_attendance_job():
    att_file = request.files.get("attendance_file")
    if not att_file:
        return err("MISSING_FILE", "请上传打卡明细文件")
    if not att_file.filename.lower().endswith((".xlsx", ".xls")):
        return err("FILE_TYPE_MISMATCH", f"仅支持 Excel 文件（.xlsx / .xls），收到：{att_file.filename}")

    shift_config = {}
    if request.form.get("shift_start"):
        shift_config["start"] = request.form["shift_start"]
    if request.form.get("shift_end"):
        shift_config["end"] = request.form["shift_end"]
    if request.form.get("ot_threshold_min"):
        shift_config["ot_threshold_min"] = int(request.form["ot_threshold_min"])

    job_id = str(uuid.uuid4())
    job_dir = UPLOAD_DIR / job_id
    job_dir.mkdir(parents=True, exist_ok=True)
    file_path = job_dir / att_file.filename
    att_file.save(str(file_path))

    try:
        result = attendance_analyze(file_path, shift_config or None)
    except ValueError as exc:
        parts = str(exc).split("|")
        code = parts[0] if len(parts) >= 2 else "ANALYZE_ERROR"
        msg = parts[1] if len(parts) >= 2 else str(exc)
        return err(code, msg, 422)

    db = get_db()
    created_at = datetime.now(tz=timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
    db.execute(
        "INSERT INTO job(id, module_key, period, summary, created_by, created_at, finished_at)"
        " VALUES(?,?,?,?,?,?,?)",
        (job_id, "attendance", None,
         json.dumps(result["summary"], ensure_ascii=False),
         g.current_user, created_at, created_at),
    )
    for f_info in result["files"]:
        db.execute(
            "INSERT INTO job_file(job_id, side, original_name, sha256_before, sha256_after, readonly_verified)"
            " VALUES(?,?,?,?,?,?)",
            (job_id, f_info["side"], att_file.filename,
             f_info["sha256_before"], f_info["sha256_after"],
             1 if f_info["readonly_verified"] else 0),
        )
    for i, row in enumerate(result["rows"]):
        db.execute(
            "INSERT INTO job_row(job_id, row_index, match_key, category, payload)"
            " VALUES(?,?,?,?,?)",
            (job_id, i, row.get("match_key"), row["category"],
             json.dumps(row, ensure_ascii=False)),
        )
    db.execute(
        "INSERT INTO audit_log(username, action, target_id, created_at) VALUES(?,?,?,?)",
        (g.current_user, "create_attendance_job", job_id, created_at),
    )
    db.commit()

    return jsonify({
        "id": job_id,
        "module_key": "attendance",
        "status": "done",
        "summary": result["summary"],
        "created_by": g.current_user,
        "created_at": created_at,
        "files": [
            {
                "side": f["side"],
                "original_name": att_file.filename,
                "sha256_before": f["sha256_before"],
                "sha256_after": f["sha256_after"],
                "readonly_verified": f["readonly_verified"],
            }
            for f in result["files"]
        ],
    }), 201


# ── 社保公积金核对 ────────────────────────────────────────────────────────
@app.route("/api/social-ins/jobs", methods=["POST"])
@require_auth
def create_social_ins_job():
    payroll_file = request.files.get("payroll_file")
    social_file = request.files.get("social_file")
    if not payroll_file or not social_file:
        return err("MISSING_FILE", "请同时上传工资表和社保申报表")
    for f in [payroll_file, social_file]:
        if not f.filename.lower().endswith((".xlsx", ".xls")):
            return err("FILE_TYPE_MISMATCH", f"仅支持 Excel 文件（.xlsx / .xls），收到：{f.filename}")

    job_id = str(uuid.uuid4())
    job_dir = UPLOAD_DIR / job_id
    job_dir.mkdir(parents=True, exist_ok=True)
    payroll_path = job_dir / f"payroll_{payroll_file.filename}"
    social_path = job_dir / f"social_{social_file.filename}"
    payroll_file.save(str(payroll_path))
    social_file.save(str(social_path))

    try:
        result = social_reconcile(payroll_path, social_path)
    except ValueError as exc:
        parts = str(exc).split("|")
        code = parts[0] if len(parts) >= 2 else "RECONCILE_ERROR"
        msg = parts[1] if len(parts) >= 2 else str(exc)
        shutil.rmtree(job_dir, ignore_errors=True)
        return err(code, msg, 422)

    db = get_db()
    created_at = now_iso()
    db.execute(
        "INSERT INTO job(id, module_key, period, summary, created_by, created_at, finished_at)"
        " VALUES(?,?,?,?,?,?,?)",
        (job_id, "social_ins", None,
         json.dumps(result["summary"], ensure_ascii=False),
         g.current_user, created_at, created_at),
    )
    for f_info in result["files"]:
        f_meta = payroll_file if f_info["side"] == "payroll" else social_file
        db.execute(
            "INSERT INTO job_file(job_id, side, original_name, sha256_before, sha256_after, readonly_verified)"
            " VALUES(?,?,?,?,?,?)",
            (job_id, f_info["side"], f_meta.filename,
             f_info["sha256_before"], f_info["sha256_after"],
             1 if f_info["readonly_verified"] else 0),
        )
    for i, row in enumerate(result["rows"]):
        db.execute(
            "INSERT INTO job_row(job_id, row_index, match_key, category, payload)"
            " VALUES(?,?,?,?,?)",
            (job_id, i, row.get("match_key"), row["category"],
             json.dumps(row, ensure_ascii=False)),
        )
    db.execute(
        "INSERT INTO audit_log(username, action, target_id, created_at) VALUES(?,?,?,?)",
        (g.current_user, "create_social_ins_job", job_id, created_at),
    )
    db.commit()

    return jsonify({
        "id": job_id,
        "module_key": "social_ins",
        "status": "done",
        "summary": result["summary"],
        "created_by": g.current_user,
        "created_at": created_at,
        "files": [
            {
                "side": f["side"],
                "original_name": (payroll_file if f["side"] == "payroll" else social_file).filename,
                "sha256_before": f["sha256_before"],
                "sha256_after": f["sha256_after"],
                "readonly_verified": f["readonly_verified"],
            }
            for f in result["files"]
        ],
    }), 201


# ── 人员成本分析 ─────────────────────────────────────────────────────────
@app.route("/api/cost-analysis/jobs", methods=["POST"])
@require_auth
def create_cost_analysis_job():
    staff_file = request.files.get("staff_file")
    if not staff_file:
        return err("MISSING_FILE", "请上传在职人员表")
    if not staff_file.filename.lower().endswith((".xlsx", ".xls")):
        return err("FILE_TYPE_MISMATCH", f"仅支持 Excel 文件（.xlsx / .xls），收到：{staff_file.filename}")

    job_id = str(uuid.uuid4())
    job_dir = UPLOAD_DIR / job_id
    job_dir.mkdir(parents=True, exist_ok=True)
    file_path = job_dir / staff_file.filename
    staff_file.save(str(file_path))

    try:
        result = cost_analyze(file_path)
    except ValueError as exc:
        parts = str(exc).split("|")
        code = parts[0] if len(parts) >= 2 else "ANALYZE_ERROR"
        msg = parts[1] if len(parts) >= 2 else str(exc)
        shutil.rmtree(job_dir, ignore_errors=True)
        return err(code, msg, 422)

    db = get_db()
    created_at = now_iso()
    db.execute(
        "INSERT INTO job(id, module_key, period, summary, created_by, created_at, finished_at)"
        " VALUES(?,?,?,?,?,?,?)",
        (job_id, "cost_analysis", None,
         json.dumps(result["summary"], ensure_ascii=False),
         g.current_user, created_at, created_at),
    )
    for f_info in result["files"]:
        db.execute(
            "INSERT INTO job_file(job_id, side, original_name, sha256_before, sha256_after, readonly_verified)"
            " VALUES(?,?,?,?,?,?)",
            (job_id, f_info["side"], staff_file.filename,
             f_info["sha256_before"], f_info["sha256_after"],
             1 if f_info["readonly_verified"] else 0),
        )
    for i, row in enumerate(result["rows"]):
        db.execute(
            "INSERT INTO job_row(job_id, row_index, match_key, category, payload)"
            " VALUES(?,?,?,?,?)",
            (job_id, i, row.get("match_key"), row["category"],
             json.dumps(row, ensure_ascii=False)),
        )
    db.execute(
        "INSERT INTO audit_log(username, action, target_id, created_at) VALUES(?,?,?,?)",
        (g.current_user, "create_cost_analysis_job", job_id, created_at),
    )
    db.commit()

    return jsonify({
        "id": job_id,
        "module_key": "cost_analysis",
        "status": "done",
        "summary": result["summary"],
        "created_by": g.current_user,
        "created_at": created_at,
        "files": [
            {
                "side": f["side"],
                "original_name": staff_file.filename,
                "sha256_before": f["sha256_before"],
                "sha256_after": f["sha256_after"],
                "readonly_verified": f["readonly_verified"],
            }
            for f in result["files"]
        ],
    }), 201


@app.route("/api/jobs", methods=["GET"])
@require_auth
def list_jobs():
    db = get_db()
    page = max(1, int(request.args.get("page", 1)))
    page_size = min(50, int(request.args.get("page_size", 20)))
    offset = (page - 1) * page_size

    module = request.args.get("module", "").strip()
    if module:
        rows = db.execute(
            "SELECT * FROM job WHERE module_key=? ORDER BY created_at DESC LIMIT ? OFFSET ?",
            (module, page_size, offset),
        ).fetchall()
        total = db.execute("SELECT COUNT(*) FROM job WHERE module_key=?", (module,)).fetchone()[0]
    else:
        rows = db.execute(
            "SELECT * FROM job ORDER BY created_at DESC LIMIT ? OFFSET ?",
            (page_size, offset),
        ).fetchall()
        total = db.execute("SELECT COUNT(*) FROM job").fetchone()[0]

    items = []
    for r in rows:
        files = db.execute("SELECT * FROM job_file WHERE job_id=?", (r["id"],)).fetchall()
        items.append({
            "id": r["id"],
            "period": r["period"],
            "status": r["status"],
            "summary": json.loads(r["summary"]) if r["summary"] else None,
            "created_by": r["created_by"],
            "created_at": r["created_at"],
            "files": [dict(f) for f in files],
        })
    return jsonify({"items": items, "total": total, "page": page, "page_size": page_size})


@app.route("/api/jobs/<job_id>", methods=["GET"])
@require_auth
def get_job(job_id):
    db = get_db()
    job = db.execute("SELECT * FROM job WHERE id=?", (job_id,)).fetchone()
    if not job:
        return err("JOB_NOT_FOUND", f"作业 {job_id} 不存在", 404)
    files = db.execute("SELECT * FROM job_file WHERE job_id=?", (job_id,)).fetchall()
    return jsonify({
        "id": job["id"],
        "period": job["period"],
        "status": job["status"],
        "summary": json.loads(job["summary"]) if job["summary"] else None,
        "created_by": job["created_by"],
        "created_at": job["created_at"],
        "files": [dict(f) for f in files],
    })


@app.route("/api/jobs/<job_id>/rows", methods=["GET"])
@require_auth
def get_job_rows(job_id):
    db = get_db()
    if not db.execute("SELECT 1 FROM job WHERE id=?", (job_id,)).fetchone():
        return err("JOB_NOT_FOUND", f"作业 {job_id} 不存在", 404)

    category = request.args.getlist("category")  # 可多值
    keyword = request.args.get("keyword", "").strip()
    page = max(1, int(request.args.get("page", 1)))
    page_size = min(200, int(request.args.get("page_size", 100)))

    q = "SELECT payload FROM job_row WHERE job_id=?"
    params: list = [job_id]
    if category:
        placeholders = ",".join("?" * len(category))
        q += f" AND category IN ({placeholders})"
        params.extend(category)
    q += " ORDER BY row_index"

    all_rows = [json.loads(r["payload"]) for r in db.execute(q, params).fetchall()]

    if keyword:
        kw = keyword.lower()
        all_rows = [
            r for r in all_rows
            if kw in (r.get("match_key") or "").lower()
            or kw in (r.get("name") or "").lower()
        ]

    total = len(all_rows)
    offset = (page - 1) * page_size
    items = all_rows[offset: offset + page_size]
    return jsonify({"items": items, "total": total, "page": page, "page_size": page_size})


@app.route("/api/jobs/<job_id>/export", methods=["GET"])
@require_auth
def export_job(job_id):
    db = get_db()
    job = db.execute("SELECT * FROM job WHERE id=?", (job_id,)).fetchone()
    if not job:
        return err("JOB_NOT_FOUND", f"作业 {job_id} 不存在", 404)

    rows = [
        json.loads(r["payload"])
        for r in db.execute(
            "SELECT payload FROM job_row WHERE job_id=? ORDER BY row_index", (job_id,)
        ).fetchall()
    ]
    summary = json.loads(job["summary"]) if job["summary"] else {}

    wb = openpyxl.Workbook()
    module_key = job["module_key"]

    if module_key == "contract_alert":
        _write_contract_sheet(wb.active, rows)
        wb.active.title = "合同到期预警"
        ws2 = wb.create_sheet("预警清单")
        warn_rows = [r for r in rows if r["category"] != "normal"]
        _write_contract_sheet(ws2, warn_rows)
    elif module_key == "roster_check":
        _write_roster_sheet(wb.active, rows)
        wb.active.title = "花名册校验报告"
        ws2 = wb.create_sheet("问题清单")
        issue_rows = [r for r in rows if r["category"] != "clean"]
        _write_roster_sheet(ws2, issue_rows)
    elif module_key == "attendance":
        summary_rows = [r for r in rows if r.get("row_type") == "person_summary"]
        detail_rows = [r for r in rows if r.get("row_type") == "detail"]
        _write_attendance_summary_sheet(wb.active, summary_rows)
        wb.active.title = "考勤汇总"
        ws2 = wb.create_sheet("异常明细")
        abnormal = [r for r in detail_rows if r["category"] != "normal"]
        _write_attendance_detail_sheet(ws2, abnormal)
        ws3 = wb.create_sheet("完整记录")
        _write_attendance_detail_sheet(ws3, detail_rows)
    elif module_key == "social_ins":
        _write_social_sheet(wb.active, rows)
        wb.active.title = "社保核对明细"
        ws2 = wb.create_sheet("差异清单")
        diff_rows = [r for r in rows if r["category"] != "match"]
        _write_social_sheet(ws2, diff_rows)
    elif module_key == "cost_analysis":
        dept_rows = [r for r in rows if r.get("row_type") == "dept_summary"]
        level_rows = [r for r in rows if r.get("row_type") == "dept_level"]
        detail_rows = [r for r in rows if r.get("row_type") == "detail"]
        _write_cost_dept_sheet(wb.active, dept_rows)
        wb.active.title = "部门汇总"
        ws2 = wb.create_sheet("部门-职级明细")
        _write_cost_level_sheet(ws2, level_rows)
        ws3 = wb.create_sheet("人员清单")
        _write_cost_detail_sheet(ws3, detail_rows)
    else:
        _write_detail_sheet(wb.active, rows, summary)
        wb.active.title = "完整对账明细"
        ws2 = wb.create_sheet("差异清单")
        diff_rows = [r for r in rows if r["category"] != "match"]
        _write_diff_sheet(ws2, diff_rows)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    period = job["period"] or "未知期间"
    if module_key == "contract_alert":
        filename = f"合同到期预警_{job['created_at'][:10]}.xlsx"
    elif module_key == "roster_check":
        filename = f"花名册校验_{job['created_at'][:10]}.xlsx"
    elif module_key == "attendance":
        filename = f"考勤统计_{job['created_at'][:10]}.xlsx"
    elif module_key == "social_ins":
        filename = f"社保核对_{job['created_at'][:10]}.xlsx"
    elif module_key == "cost_analysis":
        filename = f"人员成本分析_{job['created_at'][:10]}.xlsx"
    else:
        filename = f"薪资对账_{period}_{job['created_at'][:10]}.xlsx"

    db.execute(
        "INSERT INTO audit_log(username, action, target_id, created_at) VALUES(?,?,?,?)",
        (g.current_user, "export_result", job_id, now_iso()),
    )
    db.commit()

    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


# ── Excel 导出辅助 ────────────────────────────────────────────────────────────
CATEGORY_LABELS = {
    "match": "一致",
    "amount_diff": "金额不符",
    "payroll_only": "工资表有银行无",
    "bank_only": "银行有工资表无",
}

DETAIL_HEADERS = ["工号", "姓名", "部门", "对账结果", "工资表实发金额", "银行代发金额", "差额（银行-工资表）"]
DIFF_HEADERS = DETAIL_HEADERS + ["差异说明"]

CONTRACT_CATEGORY_LABELS = {
    "expired": "已过期",
    "within_30": "30天内到期",
    "within_60": "60天内到期",
    "within_90": "90天内到期",
    "normal": "正常",
}

CONTRACT_HEADERS = ["工号", "姓名", "部门", "合同类型", "合同开始日期", "合同结束日期", "剩余天数", "预警等级"]


def _write_detail_sheet(ws, rows: list[dict], summary: dict):
    ws.append(DETAIL_HEADERS)
    for r in rows:
        ws.append([
            r.get("match_key") or "",
            r.get("name") or "",
            r.get("dept") or "",
            CATEGORY_LABELS.get(r["category"], r["category"]),
            r.get("payroll_amount") or "—",
            r.get("bank_amount") or "—",
            r.get("delta") or "—",
        ])


def _write_diff_sheet(ws, rows: list[dict]):
    ws.append(DIFF_HEADERS)
    for r in rows:
        cat = r["category"]
        delta = r.get("delta")
        if cat == "amount_diff" and delta:
            try:
                d = float(delta)
                desc = f"银行{'少' if d < 0 else '多'}发 {abs(d):.2f} 元"
            except Exception:
                desc = delta
        elif cat == "payroll_only":
            desc = "工资表有记录，银行无代发"
        else:
            desc = "银行有代发，工资表无记录"

        ws.append([
            r.get("match_key") or "",
            r.get("name") or "",
            r.get("dept") or "",
            CATEGORY_LABELS.get(cat, cat),
            r.get("payroll_amount") or "—",
            r.get("bank_amount") or "—",
            r.get("delta") or "—",
            desc,
        ])


def _write_contract_sheet(ws, rows: list[dict]):
    ws.append(CONTRACT_HEADERS)
    for r in rows:
        ws.append([
            r.get("match_key") or "",
            r.get("name") or "",
            r.get("dept") or "",
            r.get("contract_type") or "",
            r.get("start_date") or "",
            r.get("end_date") or "",
            r.get("days_remaining") if r.get("days_remaining") is not None else "—",
            CONTRACT_CATEGORY_LABELS.get(r["category"], r["category"]),
        ])


ROSTER_CATEGORY_LABELS = {
    "clean": "正常",
    "missing_field": "缺失必填项",
    "duplicate": "重复",
    "contradiction": "矛盾",
}

ROSTER_HEADERS = ["工号", "姓名", "性别", "出生日期", "身份证号", "部门", "入职日期", "问题类型", "问题详情"]


def _write_roster_sheet(ws, rows: list[dict]):
    ws.append(ROSTER_HEADERS)
    for r in rows:
        ws.append([
            r.get("match_key") or "",
            r.get("name") or "",
            r.get("gender") or "",
            r.get("birth_date") or "",
            r.get("id_number") or "",
            r.get("dept") or "",
            r.get("hire_date") or "",
            ROSTER_CATEGORY_LABELS.get(r["category"], r["category"]),
            r.get("issue_summary") or "",
        ])


ATTENDANCE_CATEGORY_LABELS = {
    "normal": "正常",
    "late": "迟到",
    "early": "早退",
    "missing_punch": "缺卡",
    "overtime": "加班",
    "has_issue": "有异常",
}

ATTENDANCE_SUMMARY_HEADERS = ["工号", "姓名", "部门", "出勤天数", "迟到次数", "早退次数", "缺卡次数", "加班时长(h)"]
ATTENDANCE_DETAIL_HEADERS = ["日期", "工号", "姓名", "部门", "上班打卡", "下班打卡", "迟到(分)", "早退(分)", "加班(分)", "状态"]


def _write_attendance_summary_sheet(ws, rows: list[dict]):
    ws.append(ATTENDANCE_SUMMARY_HEADERS)
    for r in rows:
        ws.append([
            r.get("match_key") or "",
            r.get("name") or "",
            r.get("dept") or "",
            r.get("days", 0),
            r.get("late_count", 0),
            r.get("early_count", 0),
            r.get("missing_count", 0),
            r.get("ot_hours", 0),
        ])


def _write_attendance_detail_sheet(ws, rows: list[dict]):
    ws.append(ATTENDANCE_DETAIL_HEADERS)
    for r in rows:
        ws.append([
            r.get("att_date") or "",
            r.get("match_key") or "",
            r.get("name") or "",
            r.get("dept") or "",
            r.get("clock_in") or "—",
            r.get("clock_out") or "—",
            r.get("late_min", 0) if r.get("is_late") else "",
            r.get("early_min", 0) if r.get("is_early") else "",
            r.get("ot_min", 0) if r.get("is_overtime") else "",
            ATTENDANCE_CATEGORY_LABELS.get(r["category"], r["category"]),
        ])


SOCIAL_CATEGORY_LABELS = {
    "match": "一致",
    "base_diff": "基数差异",
    "amount_diff": "金额差异",
    "payroll_only": "工资表有社保无",
    "social_only": "社保有工资表无",
}

SOCIAL_HEADERS = [
    "工号", "姓名", "部门", "核对结果",
    "工资表社保基数", "社保表社保基数", "社保基数差",
    "工资表公积金基数", "社保表公积金基数", "公积金基数差",
    "工资表社保个人", "社保表社保个人", "社保个人差",
    "工资表社保单位", "社保表社保单位", "社保单位差",
    "工资表公积金个人", "社保表公积金个人", "公积金个人差",
    "工资表公积金单位", "社保表公积金单位", "公积金单位差",
    "差异字段",
]


def _write_social_sheet(ws, rows: list[dict]):
    ws.append(SOCIAL_HEADERS)
    for r in rows:
        ws.append([
            r.get("match_key") or "",
            r.get("name") or "",
            r.get("dept") or "",
            SOCIAL_CATEGORY_LABELS.get(r["category"], r["category"]),
            r.get("payroll_si_base") or "—",
            r.get("social_si_base") or "—",
            r.get("si_base_delta") or "",
            r.get("payroll_hf_base") or "—",
            r.get("social_hf_base") or "—",
            r.get("hf_base_delta") or "",
            r.get("payroll_si_person") or "—",
            r.get("social_si_person") or "—",
            r.get("si_person_delta") or "",
            r.get("payroll_si_company") or "—",
            r.get("social_si_company") or "—",
            r.get("si_company_delta") or "",
            r.get("payroll_hf_person") or "—",
            r.get("social_hf_person") or "—",
            r.get("hf_person_delta") or "",
            r.get("payroll_hf_company") or "—",
            r.get("social_hf_company") or "—",
            r.get("hf_company_delta") or "",
            "、".join(r.get("diff_fields") or []),
        ])


COST_DEPT_HEADERS = ["部门", "人数", "总月薪", "人均月薪", "占比(%)"]
COST_LEVEL_HEADERS = ["部门", "职级", "人数", "总月薪", "人均月薪", "占比(%)"]
COST_DETAIL_HEADERS = ["工号", "姓名", "部门", "职级", "月薪"]


def _write_cost_dept_sheet(ws, rows: list[dict]):
    ws.append(COST_DEPT_HEADERS)
    for r in rows:
        ws.append([
            r.get("dept") or "",
            r.get("headcount", 0),
            r.get("total_cost") or "",
            r.get("avg_cost") or "",
            r.get("pct") or "",
        ])


def _write_cost_level_sheet(ws, rows: list[dict]):
    ws.append(COST_LEVEL_HEADERS)
    for r in rows:
        ws.append([
            r.get("dept") or "",
            r.get("level") or "",
            r.get("headcount", 0),
            r.get("total_cost") or "",
            r.get("avg_cost") or "",
            r.get("pct") or "",
        ])


def _write_cost_detail_sheet(ws, rows: list[dict]):
    ws.append(COST_DETAIL_HEADERS)
    for r in rows:
        ws.append([
            r.get("match_key") or "",
            r.get("name") or "",
            r.get("dept") or "",
            r.get("level") or "",
            r.get("monthly_salary") or "",
        ])


def _job_response(job_id, period, summary, created_at, files_info, p_name, b_name):
    return {
        "id": job_id,
        "period": period or None,
        "status": "done",
        "summary": summary,
        "created_at": created_at,
        "files": [
            {
                "side": f["side"],
                "original_name": p_name if f["side"] == "payroll" else b_name,
                "sha256_before": f["sha256_before"],
                "sha256_after": f["sha256_after"],
                "readonly_verified": f["readonly_verified"],
            }
            for f in files_info
        ],
    }


# ── 启动 ─────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    init_db()
    print("=" * 50)
    print("HR 智能工作台 启动中...")
    print(f"访问地址: http://localhost:8080")
    print("内置账号: hr01/hr123456  fin01/fin123456  admin/admin888")
    print("=" * 50)
    app.run(host="0.0.0.0", port=8080, debug=False)
